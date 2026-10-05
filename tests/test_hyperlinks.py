import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.worksheet.hyperlink import Hyperlink

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
import excel_splitter as app


class HyperlinkSplitTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / '源表.xlsx'
        self.out = self.root / 'out'

    def tearDown(self):
        self.temp.cleanup()

    def save(self, rows, title='Sheet'):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = title
        for row in rows:
            ws.append(row)
        wb.save(self.source)
        return wb, ws

    def split(self, **kwargs):
        return app.split_workbook(self.source, None, kwargs.pop('header_rows', 1), 1, self.out, **kwargs)

    def test_hyperlink_to_other_sheet_is_kept_not_fatal(self):
        # The other sheet is untouched by the split and absent from the output, so the
        # location cannot be rewritten — but it must not stop the whole run either.
        wb, ws = self.save([['部门', '链接'], ['甲', None], ['乙', None]])
        wb.create_sheet('说明')['A1'] = '见此'
        ws['B2'].hyperlink = Hyperlink(ref='B2', location='说明!A1')
        wb.save(self.source)
        file = self.split()[0]
        wb = openpyxl.load_workbook(file)
        try:
            self.assertEqual(wb.active['B2'].hyperlink.location, '说明!A1')
        finally:
            wb.close()

    def test_same_sheet_hyperlink_follows_retained_rows(self):
        wb, ws = self.save([['部门', '链接'], ['甲', None], ['乙', None], ['甲', None]])
        ws['B2'].hyperlink = Hyperlink(ref='B2', location='Sheet!A4')
        ws['B4'].hyperlink = Hyperlink(ref='B4', location='Sheet!A3')
        wb.save(self.source)
        file = self.split()[0]
        wb = openpyxl.load_workbook(file)
        try:
            self.assertEqual(wb.active['B2'].hyperlink.location, 'Sheet!A3')
            # Source row 4 lands on row 3 of the output, and its target (row 3, 乙) was dropped.
            self.assertEqual(wb.active['B3'].hyperlink.location, '#REF!')
        finally:
            wb.close()


if __name__ == '__main__':
    unittest.main()
